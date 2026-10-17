"""Incrementally export completed Boreas evaluations to one XLSX workbook.

Runtime dependency: openpyxl (CPU only). JSON reports remain authoritative.
Managed sheet rows are replaced by model/training-seed/epoch on rerun.
"""
from __future__ import annotations
import argparse
from contextlib import contextmanager
from datetime import datetime, timezone
import hashlib
import json
import math
import os
from pathlib import Path
import re
import tempfile

BEAMS = (128, 64, 32, 16)
CLASSES = ('Car', 'Pedestrian', 'Cyclist')
SCHEMA = 'boreas_excel_v1'
DEFAULT_BOOK = Path(__file__).resolve().parents[2] / 'output/boreas_summary/boreas_benchmark_results.xlsx'
IDENTITY = ['RunKey', 'Model', 'TrainSeed', 'Epoch', 'Tag']
HEADERS = {
    'Runs': IDENTITY + ['UpdatedUTC', 'TrainSeedEvidence', 'TrainPrecision', 'TrainBatchPerGPU',
        'TrainGPUs', 'TrainWorkers', 'TrainEpochs', 'EvalSeed', 'EvalPrecision', 'EvalBatch',
        'Car_AP40_pct', 'Pedestrian_AP40_pct', 'Cyclist_AP40_pct', 'mAP40_128_pct',
        'mAP40_64_pct', 'mAP40_32_pct', 'mAP40_16_pct', 'Drop128to64_pp', 'Drop128to32_pp',
        'Drop128to16_pp', 'FPS_serial', 'E2E_mean_ms', 'E2E_p50_ms', 'E2E_p95_ms',
        'Preprocess_mean_ms', 'H2D_mean_ms', 'Forward_with_NMS_mean_ms', 'Parameters',
        'CUDA_allocated_MiB', 'CUDA_reserved_MiB', 'ActiveVoxelsMedian', 'ProfileFrames',
        'WarmupFrames', 'GPU', 'PyTorch', 'CUDA', 'TF32_matmul', 'TF32_cudnn',
        'CheckpointSHA256', 'ModelConfigSHA256', 'TruthSHA256', 'APProtocol',
        'ExtendedProtocol', 'RunDirectory', 'TrainingSource', 'Notes'],
    'Accuracy': IDENTITY + ['Beam', 'Range', 'Class', 'AP40_pct', 'mAP40_pct', 'GT',
        'Predictions', 'IgnoredMisc', 'ScoreThreshold', 'TP', 'FP', 'FN', 'Precision',
        'Recall', 'F1', 'SourceJSON'],
    'Calibration': IDENTITY + ['Beam', 'Class', 'ECE', 'Brier', 'Detections',
        'MeanConfidence', 'EmpiricalPrecision', 'ScoreThreshold', 'GT', 'TP', 'FP',
        'FN', 'Recall', 'CalibrationBins', 'SourceJSON'],
    'Density': IDENTITY + ['Beam', 'Class', 'Grouping', 'Range', 'PointsBin', 'GT',
        'TP', 'FN', 'Recall', 'ScoreThreshold', 'SourceJSON'],
    'Reliability': IDENTITY + ['Beam', 'Class', 'Bin', 'Lower', 'Upper', 'Count',
        'MeanConfidence', 'EmpiricalPrecision', 'SourceJSON'],
}
NOTES = {
    'Runs': 'One completed run per model / training seed / epoch; reruns replace it. AP is percent, drops are percentage points. Seed evidence is recorded or inferred_legacy.',
    'Accuracy': 'Local rotated 3D AP40, not BEV AP. AP is percent; precision/recall/F1 are fractions. Range bins independently filter GT and prediction centers.',
    'Calibration': 'ECE and Brier are fractions on nonignored post-NMS detections; missed objects are reflected in recall. Empty populations remain blank.',
    'Density': 'Full-domain matches first. density_fixed128 uses dense GT point cohorts; density_actual uses the evaluated beam subset. Recall is a fraction.',
    'Reliability': 'Equal-width confidence bins; MeanConfidence and EmpiricalPrecision are fractions. Empty-bin means remain blank.',
}


def read_json(path):
    with Path(path).open() as stream:
        return json.load(stream)


def file_hash(path):
    digest = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(block)
    return digest.hexdigest()


def training_details(folder, metadata=None):
    files = [Path(metadata)] if metadata else sorted(folder.glob('run_config_*.json'))
    if files:
        manifests = [read_json(p) for p in files]
        identities = {(m.get('seed'), m.get('training_precision'), m.get('batch_size_per_gpu'),
                       m.get('total_gpus'), m.get('epochs'),
                       json.dumps(m.get('effective_config'), sort_keys=True)) for m in manifests}
        if len(identities) != 1:
            raise ValueError('Mixed training metadata; choose --metadata explicitly')
        m = manifests[-1]
        if not isinstance(m.get('seed'), int) or isinstance(m['seed'], bool):
            raise ValueError('Training seed missing; never infer it from a folder name')
        return {'TrainSeed': m['seed'], 'TrainSeedEvidence': 'recorded',
                'TrainPrecision': m.get('training_precision'),
                'TrainBatchPerGPU': m.get('batch_size_per_gpu'), 'TrainGPUs': m.get('total_gpus'),
                'TrainWorkers': m.get('arguments', {}).get('workers'), 'TrainEpochs': m.get('epochs'),
                'TrainingSource': str(files[-1].resolve()), 'Notes': None}
    note_path = folder / 'legacy_run_rename.json'
    if not note_path.exists():
        raise ValueError('No training metadata or documented legacy seed provenance')
    note = read_json(note_path)
    seed = note.get('historical_training_seed')
    if not isinstance(seed, int) or isinstance(seed, bool) or not note.get('seed_basis'):
        raise ValueError('Invalid legacy seed note')
    logs = [p for p in sorted(folder.glob('train_*.log'))
            if 'End training' in p.read_text(errors='replace')]
    if not logs:
        raise ValueError('Legacy run has no completed training log')
    text = logs[-1].read_text(errors='replace')
    values = {}
    for key in ('batch_size', 'workers', 'epochs', 'use_amp', 'fix_random_seed'):
        match = re.search(r'\bINFO\s+' + key + r'\s+(\S+)', text)
        values[key] = match.group(1) if match else None
    if values['fix_random_seed'] != 'True':
        raise ValueError('Legacy log does not confirm fixed-seed training')
    return {'TrainSeed': seed, 'TrainSeedEvidence': 'inferred_legacy',
            'TrainPrecision': 'AMP' if values['use_amp'] == 'True' else 'FP32',
            'TrainBatchPerGPU': int(values['batch_size']) if values['batch_size'] else None,
            'TrainGPUs': 1 if 'Training with a single process' in text else None,
            'TrainWorkers': int(values['workers']) if values['workers'] else None,
            'TrainEpochs': int(values['epochs']) if values['epochs'] else None,
            'TrainingSource': str(note_path.resolve()) + '; ' + str(logs[-1].resolve()),
            'Notes': note['seed_basis'] + ' Workbook inclusion is not certification by study.py.'}


def collect_run(root, model, tag, epoch, metadata=None):
    root = Path(root).resolve()
    folder = root / model / tag
    if not folder.resolve().is_relative_to(root) or not folder.is_dir():
        raise ValueError('Run directory missing or outside --root')
    if metadata and not Path(metadata).is_absolute():
        metadata = folder / metadata
    training = training_details(folder, metadata)
    key = f'{model}|{training["TrainSeed"]}|{epoch}'
    base = dict(zip(IDENTITY, (key, model, training['TrainSeed'], epoch, tag)))
    checkpoint = folder / 'ckpt' / f'checkpoint_epoch_{epoch}.pth'
    checkpoint_hash = file_hash(checkpoint)
    sheets = {name: [] for name in HEADERS}
    aps = {}
    common = None
    sources = []
    latest_report_ns = 0
    for beam in BEAMS:
        result = folder / 'eval' / f'epoch_{epoch}' / 'val' / f'beam{beam}'
        ap_path = result / 'range_ap.json'
        ext_path = folder / f'extended_{beam}.json'
        ap, extra = read_json(ap_path), read_json(ext_path)
        latest_report_ns = max(latest_report_ns, ap_path.stat().st_mtime_ns, ext_path.stat().st_mtime_ns)
        inf = extra.get('inference') or {}
        if ap.get('protocol_version') != 'local_ap40_v1' or extra.get('protocol_version') != 'density_calibration_v1':
            raise ValueError('Outdated metric protocol; run run_metrics.sh again')
        if extra.get('beam') != beam or not extra.get('verified_inference'):
            raise ValueError('Missing verified inference or wrong beam')
        if inf.get('checkpoint_sha256') != checkpoint_hash or Path(inf.get('checkpoint', '')).resolve() != checkpoint.resolve():
            raise ValueError('Requested checkpoint differs from evaluation weights')
        if (inf.get('precision'), inf.get('batch_size'), inf.get('score_cutoff'),
                inf.get('nms_pre_maxsize'), inf.get('nms_post_maxsize')) != ('FP32', 1, .001, 4096, 500):
            raise ValueError('Primary export requires FP32, batch 1, cutoff .001, NMS caps 4096/500')
        if not ap.get('prediction_sha256') or ap['prediction_sha256'] != inf.get('prediction_sha256'):
            raise ValueError('AP and extended prediction hashes differ')
        if not ap.get('truth_sha256') or ap['truth_sha256'] != extra.get('truth_sha256'):
            raise ValueError('AP and extended GT hashes differ')
        predicted = result / 'result.pkl'
        if (Path(ap['prediction_file']).resolve() != predicted.resolve()
                or Path(extra['prediction_file']).resolve() != predicted.resolve()
                or file_hash(predicted) != inf['prediction_sha256']):
            raise ValueError('Prediction file changed or is from a different epoch')
        identity = (extra['truth_sha256'], inf['model_config_sha256'], inf.get('evaluation_seed'),
                    inf.get('tf32_matmul'), inf.get('tf32_cudnn'), inf.get('torch_version'),
                    inf.get('cuda_version'), extra['calibration_bins'], extra['operating_score'])
        if common is not None and common != identity:
            raise ValueError('Beam reports do not share GT/config/evaluation conditions')
        common = identity
        aps[beam] = ap['ranges']['all']
        if aps[beam].get('mAP40') is None:
            raise ValueError('Missing mAP; export requires complete primary metrics')
        sources.append(str(ap_path.resolve()))
        for band, entry in ap['ranges'].items():
            for cls in CLASSES:
                item = entry['classes'][cls]
                op = item['operating_point']
                sheets['Accuracy'].append({**base, 'Beam': beam, 'Range': band, 'Class': cls,
                    'AP40_pct': item['AP40'], 'mAP40_pct': entry['mAP40'], 'GT': item['gt'],
                    'Predictions': item['predictions'], 'IgnoredMisc': item['ignored_misc'],
                    'ScoreThreshold': op['score_threshold'], **{k: op[k] for k in ('TP', 'FP', 'FN')},
                    'Precision': op['precision'], 'Recall': op['recall'], 'F1': op['F1'],
                    'SourceJSON': str(ap_path.resolve())})
        for cls in CLASSES:
            entry = extra['classes'][cls]
            cal, op = entry['calibration'], entry['operating_point']
            sheets['Calibration'].append({**base, 'Beam': beam, 'Class': cls,
                'ECE': cal['ECE'], 'Brier': cal['Brier'], 'Detections': cal['detections'],
                'MeanConfidence': cal['mean_confidence'], 'EmpiricalPrecision': cal['empirical_precision'],
                'ScoreThreshold': op['score_threshold'], 'GT': op['gt'],
                **{k: op[k] for k in ('TP', 'FP', 'FN')}, 'Recall': op['recall'],
                'CalibrationBins': extra['calibration_bins'], 'SourceJSON': str(ext_path.resolve())})
            for mode in ('density_fixed128', 'density_actual'):
                for band, groups in entry[mode].items():
                    for group, counts in groups.items():
                        sheets['Density'].append({**base, 'Beam': beam, 'Class': cls,
                            'Grouping': mode, 'Range': band, 'PointsBin': group,
                            'GT': counts['gt'], 'TP': counts['TP'], 'FN': counts['FN'],
                            'Recall': counts['recall'], 'ScoreThreshold': extra['operating_score'],
                            'SourceJSON': str(ext_path.resolve())})
            for number, item in enumerate(cal['bins'], 1):
                sheets['Reliability'].append({**base, 'Beam': beam, 'Class': cls, 'Bin': number,
                    'Lower': item['lower'], 'Upper': item['upper'], 'Count': item['count'],
                    'MeanConfidence': item['mean_confidence'], 'EmpiricalPrecision': item['empirical_precision'],
                    'SourceJSON': str(ext_path.resolve())})
    profile_path = folder / 'profile.json'
    profile = read_json(profile_path)
    if profile_path.stat().st_mtime_ns < latest_report_ns:
        raise ValueError('Profiling predates the current evaluation; finish run_metrics.sh first')
    if (profile.get('model'), profile.get('precision'), profile.get('batch_size'), profile.get('seed')) != (model, 'FP32', 1, common[2]):
        raise ValueError('Profile model/precision/batch/evaluation seed mismatch')
    if Path(profile['checkpoint']).resolve() != checkpoint.resolve():
        raise ValueError('Profile checkpoint does not match this run')
    if Path(profile['data_path']).resolve() != Path(extra['reference_data']).resolve():
        raise ValueError('Profile must use the dense reference dataset')
    if (profile.get('tf32_matmul'), profile.get('tf32_cudnn'), profile.get('torch_version'),
            profile.get('cuda_version')) != common[3:7]:
        raise ValueError('Profile/inference environment mismatch')
    run = {**base, **training, 'UpdatedUTC': datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S'),
        'EvalSeed': common[2], 'EvalPrecision': 'FP32', 'EvalBatch': 1,
        **{f'{cls}_AP40_pct': aps[128]['classes'][cls]['AP40'] for cls in CLASSES},
        **{f'mAP40_{beam}_pct': aps[beam]['mAP40'] for beam in BEAMS},
        **{f'Drop128to{beam}_pp': aps[128]['mAP40'] - aps[beam]['mAP40'] for beam in (64, 32, 16)},
        'FPS_serial': profile['fps_from_mean_end_to_end'],
        **{f'E2E_{stat}_ms': profile['end_to_end'][stat+'_ms'] for stat in ('mean', 'p50', 'p95')},
        'Preprocess_mean_ms': profile['preprocess_including_npy_io']['mean_ms'],
        'H2D_mean_ms': profile['host_to_gpu']['mean_ms'],
        'Forward_with_NMS_mean_ms': profile['model_forward_including_postprocess']['mean_ms'],
        'Parameters': profile['parameters'], 'CUDA_allocated_MiB': profile['peak_cuda_allocated_mib'],
        'CUDA_reserved_MiB': profile['peak_cuda_reserved_mib'], 'ActiveVoxelsMedian': profile['active_voxels_median'],
        'ProfileFrames': profile['frames'], 'WarmupFrames': profile['warmup_frames'], 'GPU': profile['gpu'],
        'PyTorch': profile['torch_version'], 'CUDA': profile['cuda_version'], 'TF32_matmul': common[3],
        'TF32_cudnn': common[4], 'CheckpointSHA256': checkpoint_hash, 'ModelConfigSHA256': common[1],
        'TruthSHA256': common[0], 'APProtocol': ap['protocol_version'],
        'ExtendedProtocol': extra['protocol_version'], 'RunDirectory': str(folder.resolve())}
    if any(aps[beam]['mAP40'] is None for beam in BEAMS):
        raise ValueError('Missing mAP; export requires complete primary metrics')
    sheets['Runs'].append(run)
    return key, sheets


def load_excel_library():
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError('Excel export needs openpyxl. Install in the eval environment: python -m pip install openpyxl') from exc
    return openpyxl


@contextmanager
def workbook_lock(path):
    import fcntl  # The benchmark runs in Linux/WSL; lock prevents lost concurrent writes.
    with path.with_suffix(path.suffix + '.lock').open('a') as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        yield
        fcntl.flock(lock, fcntl.LOCK_UN)


def initialize_sheets(book):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for name, columns in HEADERS.items():
        if name in book.sheetnames:
            sheet = book[name]
            existing = [c.value for c in sheet[3]]
            if existing != columns:
                raise ValueError(f'Unexpected {name} schema; refusing to overwrite workbook')
            continue
        sheet = book.create_sheet(name)
        sheet['A1'] = 'Boreas benchmark — ' + name
        sheet['A1'].font = Font(name='Arial', size=14, bold=True, color='17365D')
        sheet['A2'] = NOTES[name]
        sheet['A2'].font = Font(name='Arial', size=10, italic=True, color='595959')
        sheet.row_dimensions[3].height = 42
        for index, column in enumerate(columns, 1):
            cell = sheet.cell(3, index, column)
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            width = 18
            if column in ('Model', 'Class', 'Tag', 'Grouping'): width = 24
            if column in ('RunKey', 'GPU'): width = 30
            if column in ('UpdatedUTC', 'TrainSeedEvidence'): width = 25
            if column in ('Notes', 'TrainingSource', 'RunDirectory', 'SourceJSON'): width = 65
            if column.endswith('SHA256'): width = 70
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = 'F4'
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = '17365D' if name == 'Runs' else '7890A8'


def literal(value):
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Nonfinite numeric metric')
    if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def upsert_workbook(path, key=None, rows=None):
    openpyxl = load_excel_library()
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    path = Path(path).resolve()
    path.parent.mkdir(parents=True, exist_ok=True)
    with workbook_lock(path):
        book = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()
        if not path.exists(): book.remove(book.active)
        initialize_sheets(book)
        if key is not None:
            for name, columns in HEADERS.items():
                sheet = book[name]
                for index in range(sheet.max_row, 3, -1):
                    if sheet.cell(index, 1).value == key:
                        sheet.delete_rows(index)
                for entry in rows[name]:
                    sheet.append([literal(entry.get(column)) for column in columns])
                for cells in sheet.iter_rows(min_row=4):
                    for column, cell in zip(columns, cells):
                        cell.font = Font(name='Arial', size=10, color='202020')
                        cell.alignment = Alignment(vertical='center', horizontal='right' if isinstance(cell.value, (int,float)) else 'left')
                        if isinstance(cell.value, float):
                            cell.number_format = '0.0000' if column in ('ECE','Brier') else '0.00'
                        if column in ('Notes', 'TrainingSource', 'RunDirectory', 'SourceJSON') and isinstance(cell.value, str):
                            cell.alignment = Alignment(vertical='center', horizontal='left', wrap_text=True)
                            sheet.row_dimensions[cell.row].height = max(sheet.row_dimensions[cell.row].height or 15, 15 * math.ceil(len(cell.value) / 75))
                        elif isinstance(cell.value, int) and not isinstance(cell.value, bool):
                            cell.number_format = '#,##0' if column in ('Parameters','GT','TP','FP','FN','Count','Detections','Predictions','IgnoredMisc','ActiveVoxelsMedian') else '0'
                sheet.auto_filter.ref = f'A3:{get_column_letter(len(columns))}{max(3,sheet.max_row)}'
        descriptor, temporary = tempfile.mkstemp(prefix='.'+path.stem+'-', suffix='.xlsx', dir=path.parent)
        os.close(descriptor)
        try:
            book.save(temporary)
            book.close()
            with open(temporary, 'rb') as stream: os.fsync(stream.fileno())
            os.replace(temporary, path)
        finally:
            if os.path.exists(temporary): os.unlink(temporary)
    return path


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--root', type=Path, default=Path(__file__).resolve().parents[2]/'output/boreas_models')
    parser.add_argument('--model')
    parser.add_argument('--tag')
    parser.add_argument('--epoch', type=int, default=80)
    parser.add_argument('--metadata', type=Path)
    parser.add_argument('--workbook', type=Path, default=DEFAULT_BOOK)
    parser.add_argument('--init', action='store_true', help='Create an empty workbook without inventing results')
    parser.add_argument('--check-dependency', action='store_true')
    args = parser.parse_args()
    try:
        if args.check_dependency:
            load_excel_library(); print('Excel dependency ready (openpyxl)'); return
        if args.init:
            print('Excel workbook ready:', upsert_workbook(args.workbook)); return
        if not args.model or not args.tag:
            parser.error('--model and --tag are required unless --init or --check-dependency')
        key, rows = collect_run(args.root, args.model, args.tag, args.epoch, args.metadata)
        path = upsert_workbook(args.workbook, key, rows)
        print(f'Excel updated: {path}\nRun: {key}; seed evidence: {rows["Runs"][0]["TrainSeedEvidence"]}')
    except (OSError, ValueError, KeyError, RuntimeError) as exc:
        parser.exit(1, f'Excel export failed: {exc}\nJSON reports are preserved. After fixing the issue, rerun only export_excel.py. Close Excel if it locks the workbook.\n')


if __name__ == '__main__':
    main()
